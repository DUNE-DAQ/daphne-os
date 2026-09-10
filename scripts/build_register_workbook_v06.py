#!/usr/bin/env python3
"""Create the v0.6 server audit from the untouched, hash-pinned v0.5 workbook.

Requires openpyxl==3.1.5. Only Server Platform is reassessed; other domains are
retained as historical v0.5 evidence. Resolved means implementation, not full
hardware qualification, consortium approval or safety certification.
"""
import argparse
from collections import Counter
from copy import copy
import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

V05_SHA = '7c58f7f469523b7dd69ff3836f43d1a59bffdae49e2925bb328ac182122d8fd8'
COLOURS = {'implemented':'D9EAD3','partial':'FFF2CC','missing':'F4CCCC','contract-pending':'D9D2E9'}
LABELS = {'implemented':'Resolved implementation (scoped)','partial':'Open — partial implementation',
          'missing':'Open — missing implementation','contract-pending':'Open — contract required'}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--v05',type=Path,required=True)
    p.add_argument('--audit-csv',type=Path,required=True)
    p.add_argument('--audit-json',type=Path,required=True)
    p.add_argument('--deployment-proof',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    args = p.parse_args()
    assert not args.output.exists() and args.output.resolve()!=args.v05.resolve(), 'Never overwrite v0.5 or an existing release'
    assert sha(args.v05)==V05_SHA, 'Wrong original workbook'
    with args.audit_csv.open(newline='') as f: rows=list(csv.DictReader(f))
    meta=json.loads(args.audit_json.read_text()); deployment=json.loads(args.deployment_proof.read_text())
    assert len(rows)==251 and len({r['id'] for r in rows})==251
    assert dict(Counter(r['implementation'] for r in rows))==meta['expected_counts']
    assert deployment['success'] and deployment['candidate_deployed']
    assert deployment['source_commit']==meta['server_commit']
    wb=load_workbook(args.v05)
    original_sheets=tuple(wb.sheetnames)
    historical={s.title:[[c.value for c in row] for row in s] for s in wb if s.title not in ('Server Platform','Start Here')}
    server=wb['Server Platform']
    assert server.max_row==256 and server.max_column==30
    server['A1']='Server Platform | v0.6 | 251 entries'
    server['A2']='v0.6: green clears the implementation issue only. Qualification limits remain in AG. Other domains retain v0.5 assessments; see v0.6 Release Notes.'
    server.conditional_formatting._cf_rules.clear() # Replace the old v0.5 status-colour rules.
    added=['v0.6 implementation status','Open implementation action','Qualification / evidence limits','Current producer mapping']
    for column,title in enumerate(added,31):
        cell=server.cell(5,column,title); cell._style=copy(server['AD5']._style)
        server.column_dimensions[cell.column_letter].width=55 if column!=31 else 36
    for r in rows:
        index=int(r['workbook_row']); state=r['implementation']; group=meta['evidence_groups'][r['evidence_group']]
        assert server.cell(index,1).value==r['id'] and server.cell(index,2).value==r['variable']
        assert server.cell(index,3).value==r['meaning'] and server.cell(index,7).value==r['mode']
        server.cell(index,8,LABELS[state]+'; '+r['mapping'])
        server.cell(index,29,LABELS[state]); server.cell(index,30,r['mapping'])
        server.cell(index,31,LABELS[state])
        # Clear resolved implementation actions, but retain qualification work.
        server.cell(index,32,None if state=='implemented' else r['remaining_work'])
        server.cell(index,33,r['remaining_work']+' Evidence scope: '+group['evidence_scope']+'. '+group['limit'])
        server.cell(index,34,r['mapping'])
        for cell in server[index]:
            cell.fill=PatternFill('solid',fgColor=COLOURS[state]); cell.alignment=Alignment(vertical='top',wrap_text=True)
        server.row_dimensions[index].height=72
    table=server.tables['T_Server_Platform']
    table.ref='A5:AH256'
    from openpyxl.worksheet.table import TableColumn
    for i,title in enumerate(added,31): table.tableColumns.append(TableColumn(id=i,name=title))
    table.autoFilter.ref=table.ref
    server.auto_filter.ref=table.ref; server.print_area='A1:AH256'
    wb['Start Here']['B4']='Operations variable ownership — v0.6 server audit; non-server domains retained from v0.5'
    release=wb.create_sheet('v0.6 Release Notes',0)
    release.append(['DAPHNE register workbook v0.6','DAPHNE-015 server checkpoint'])
    notes=[('Deployed source',deployment['source_commit']),('Executable SHA-256',deployment['candidate_sha256']),
       ('Firmware','Self-trigger 3f17f1b / ABI 2.0, unchanged'),
       ('Resolved scope','Server implementation issues only. Green is not both-mode hardware qualification or an overall healthy-board claim.'),
       ('Open work','Server Open Issues contains unresolved implementation rows. Server Resolved retains outstanding qualification for implemented rows.'),
       ('Other tabs','Non-server domains and cross-domain summaries retain v0.5 source assessments; they have NOT been silently cleared or requalified.'),
       ('SC boundary','SC003 owns the requested BiasEnable state. I288 is server readback. DAQ Configure preserves the enable bit; authenticated SC command ownership remains open.'),
       ('Missing','Calibrated current/analog verification; fitted mezzanine and Qt tests; approved timing/MAC/Hermes mapping; three SFP paths; physical fan conversion/stall policy; bus counters; broader command audit/SC authorization; routed firmware/full-stream; boot/network recovery.'),
       ('Original workbook SHA-256',V05_SHA),('Audit CSV SHA-256',sha(args.audit_csv)),
       ('Audit metadata SHA-256',sha(args.audit_json)),('Deployment evidence SHA-256',sha(args.deployment_proof))]
    for state,count in meta['expected_counts'].items(): notes.append((state,count))
    for row in notes: release.append(row)
    release.column_dimensions['A'].width=32; release.column_dimensions['B'].width=110
    for row in release:
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
        release.row_dimensions[row[0].row].height=45
    release.freeze_panes='B2'
    headers=['ID','Variable','Implementation','Current mapping','Open implementation action','Qualification still open','Evidence scope']
    for title,resolved in [('Server Open Issues',False),('Server Resolved',True)]:
        sheet=wb.create_sheet(title)
        sheet.append(headers)
        selected=[r for r in rows if (r['implementation']=='implemented')==resolved]
        for r in selected:
            group=meta['evidence_groups'][r['evidence_group']]
            sheet.append([r['id'],r['variable'],LABELS[r['implementation']],r['mapping'],
                          None if resolved else r['remaining_work'],r['remaining_work'],group['evidence_scope']+'; '+group['limit']])
            for cell in sheet[sheet.max_row]:
                cell.fill=PatternFill('solid',fgColor=COLOURS[r['implementation']]); cell.alignment=Alignment(vertical='top',wrap_text=True)
            sheet.row_dimensions[sheet.max_row].height=85
        for col,width in zip('ABCDEFG',[12,38,34,70,75,75,85]): sheet.column_dimensions[col].width=width
        summary_table=Table(displayName=title.replace(' ','_'),ref=f'A1:G{sheet.max_row}')
        summary_table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False)
        sheet.add_table(summary_table); sheet.freeze_panes='C2'; sheet.sheet_view.zoomScale=80
    for sheet in (release,wb['Server Open Issues'],wb['Server Resolved']):
        for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor='17365D'); cell.font=Font(color='FFFFFF',bold=True)
    wb.active=0
    wb.properties.title='DAPHNE Operations Variable Ownership Draft v0.6'
    wb.properties.description='Server implementation closure checkpoint; physical qualification and non-server domains remain explicitly scoped.'
    args.output.parent.mkdir(parents=True,exist_ok=True)
    wb.save(args.output)
    check=load_workbook(args.output)
    assert set(original_sheets)<=set(check.sheetnames) and len(check.sheetnames)==29
    for title,values in historical.items(): assert [[c.value for c in row] for row in check[title]]==values
    assert check['Server Resolved'].max_row-1==meta['expected_counts']['implemented']
    assert check['Server Open Issues'].max_row-1==251-meta['expected_counts']['implemented']
    for r in rows:
        idx=int(r['workbook_row']); assert check['Server Platform'].cell(idx,1).value==r['id']
        assert (check['Server Platform'].cell(idx,32).value is None)==(r['implementation']=='implemented')
        assert check['Server Platform'].cell(idx,33).value
    assert sha(args.v05)==V05_SHA
    print(json.dumps({'success':True,'version':'0.6','worksheets':29,'server_rows':251,
        'resolved_implementation':meta['expected_counts']['implemented'],
        'open_implementation':251-meta['expected_counts']['implemented'],
        'original_v05_unchanged':True,'non_server_cell_values_preserved':True,
        'workbook_sha256':sha(args.output),'source_commit':deployment['source_commit']},indent=2))


if __name__=='__main__': main()
